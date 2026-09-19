import calendar
import datetime
import io
import random
import openpyxl
from openpyxl.styles import Border, Side, PatternFill
import pandas as pd
import streamlit as st

# ページ基本設定
st.set_page_config(page_title="シフト自動作成システム", layout="wide")

st.title("🗓️ 葬儀支社向け シフト自動作成システム")
st.caption(
    "全支社共通ルール（社員最低2名出勤・最大5連勤制限）および支社別ルール（顧客重複防止・遅番自動分散＆翌日休み・健診公休・希望出勤対応）"
)

# ---------------------------------------------------------
# 1. 六曜計算ロジック
# ---------------------------------------------------------
NEW_MOONS = [
    (datetime.date(2025, 12, 19), 2025, 11),
    (datetime.date(2026, 1, 19), 2025, 12),
    (datetime.date(2026, 2, 17), 2026, 1),
    (datetime.date(2026, 3, 19), 2026, 2),
    (datetime.date(2026, 4, 17), 2026, 3),
    (datetime.date(2026, 5, 17), 2026, 4),
    (datetime.date(2026, 6, 15), 2026, 5),
    (datetime.date(2026, 7, 14), 2026, 6),
    (datetime.date(2026, 8, 13), 2026, 7),
    (datetime.date(2026, 9, 11), 2026, 8),
    (datetime.date(2026, 10, 10), 2026, 9),
    (datetime.date(2026, 11, 9), 2026, 10),
    (datetime.date(2026, 12, 9), 2026, 11),
    (datetime.date(2027, 1, 7), 2026, 12),
]
ROKUYO_SHORT = ["大", "赤", "勝", "友", "負", "仏"]
WEEKDAYS_JP = ["月", "火", "水", "木", "金", "土", "日"]


def get_rokuyo_short(dt: datetime.date) -> str:
  last_nm = None
  for nm_date, l_year, l_month in NEW_MOONS:
    if dt >= nm_date:
      last_nm = (nm_date, l_year, l_month)
    else:
      break
  if not last_nm:
    return ""
  nm_date, l_year, l_month = last_nm
  l_day = (dt - nm_date).days + 1
  return ROKUYO_SHORT[(l_month + l_day) % 6]


# ---------------------------------------------------------
# 2. タブ画面の構築
# ---------------------------------------------------------
tab1, tab2, tab3 = st.tabs(
    ["⚙️ 1. スタッフ・支社ルール設定", "📅 2. 対象年月・事前予定入力", "🚀 3. シフト自動作成・出力"]
)

# --- タブ1：マスタ管理 ---
with tab1:
  st.subheader("支社スタッフ情報・制約ルールの登録")
  st.write(
      "担当顧客が同じスタッフ同士の休み重複防止や、遅番対応の可否を設定します。"
  )

  default_staff = pd.DataFrame([
      {
          "氏名": "齊田",
          "区分": "社員",
          "公休数": 9,
          "担当顧客グループ": "グループA",
          "遅番対応": True,
      },
      {
          "氏名": "久保",
          "区分": "社員",
          "公休数": 9,
          "担当顧客グループ": "グループA",
          "遅番対応": True,
      },
      {
          "氏名": "上田",
          "区分": "社員",
          "公休数": 10,
          "担当顧客グループ": "グループB",
          "遅番対応": True,
      },
      {
          "氏名": "石山",
          "区分": "社員",
          "公休数": 10,
          "担当顧客グループ": "グループB",
          "遅番対応": False,
      },
      {
          "氏名": "森川",
          "区分": "パート",
          "公休数": 19,
          "担当顧客グループ": "共通",
          "遅番対応": False,
      },
      {
          "氏名": "宮地",
          "区分": "パート",
          "公休数": 9,
          "担当顧客グループ": "共通",
          "遅番対応": False,
      },
      {
          "氏名": "鈴村",
          "区分": "パート",
          "公休数": 18,
          "担当顧客グループ": "共通",
          "遅番対応": False,
      },
  ])

  edited_staff = st.data_editor(
      default_staff, num_rows="dynamic", use_container_width=True
  )

# --- タブ2：対象年月選択 ＆ 事前予定入力 ---
with tab2:
  st.subheader("シフト作成対象月 ＆ 事前予定入力")
  col_y, col_m = st.columns(2)
  with col_y:
    target_year = st.number_input("作成年", value=2026, step=1)
  with col_m:
    target_month = st.number_input(
        "作成月", value=11, min_value=1, max_value=12, step=1
    )

  _, num_days = calendar.monthrange(target_year, target_month)
  days = list(range(1, num_days + 1))
  st.info(f"対象期間: {target_year}年{target_month}月1日 ～ {num_days}日")

  tomobiki_list = []
  for d in days:
    dt = datetime.date(target_year, target_month, d)
    if get_rokuyo_short(dt) == "友":
      tomobiki_list.append(f"{d}日({WEEKDAYS_JP[dt.weekday()]})")
  st.write(f"**当月の友引日:** {', '.join(tomobiki_list)}")

  st.markdown("---")
  st.subheader("📝 事前予定（希望休・希望出勤・有給・健康診断など）の入力")
  st.caption(
      "※指定がある日のみ入力してください。"
      "【入力記号ルール】 `休` = 希望休 | `出` = 希望出勤日 | `有` = 有給休暇 | `健` = 健康診断  "
      "※ `出`（希望出勤日）に指定した日は、公休や遅番などの自動割り当てから除外され、必ず出勤となります。"
  )

  # 初期データの作成（スタッフ名 × 1日〜31日の空表）
  pre_input_data = {"氏名": edited_staff["氏名"].tolist()}
  for d in days:
    pre_input_data[f"{d}日"] = ""

  initial_pre_df = pd.DataFrame(pre_input_data)

  # 画面上で希望を入力できるテーブル
  edited_pre_df = st.data_editor(
      initial_pre_df, use_container_width=True, hide_index=True
  )

# --- タブ3：自動作成＆ダウンロード ---
with tab3:
  st.subheader("たたき台シフトの自動生成")
  st.write("ボタンをクリックすると、設定ルールに基づいて自動計算を行います。")

  if st.button("🚀 シフト自動作成を実行する", type="primary"):
    with st.spinner("ルールを満たす最適な組み合わせを計算中..."):
      try:
        wb = openpyxl.load_workbook("シフト（自動化）.xlsx")
        ws = wb["原本"]
      except FileNotFoundError:
        st.error(
            "【エラー】'シフト（自動化）.xlsx' が見つかりません。GitHubリポジトリ内に配置されているか確認してください。"
        )
        st.stop()

      # プルダウン（入力規制）クリア
      ws.data_validations.dataValidation.clear()

      # タイトル・日付・曜日・六曜のセット
      ws.cell(
          row=1, column=3
      ).value = (
          f"          {target_year}年{target_month}月≪休日10日≫ 休日シフト表"
      )
      start_date = datetime.date(target_year, target_month, 1)
      ws.cell(row=7, column=3).value = start_date

      tomobiki_days = []
      thick_side = Side(style='thick')  # くっきりした極太線
      # 友引の色（薄緑）と友引前日の色（さらに1段薄い緑）
      tomobiki_fill = PatternFill(
          start_color='E2EFDA', end_color='E2EFDA', fill_type='solid'
      )
      zenjitsu_fill = PatternFill(
          start_color='F2F9EC', end_color='F2F9EC', fill_type='solid'
      )

      for day in range(1, 32):
        col_idx = 3 + day
        if day <= num_days:
          dt = datetime.date(target_year, target_month, day)
          ws.cell(row=7, column=col_idx).value = dt
          ws.cell(row=8, column=col_idx).value = WEEKDAYS_JP[dt.weekday()]

          # 当日の六曜と翌日の六曜を取得
          rokuyo = get_rokuyo_short(dt)
          next_dt = dt + datetime.timedelta(days=1)
          next_rokuyo = get_rokuyo_short(next_dt)

          # 友引・友引前日の表記 ＆ 着色（5,6行目は除外し、7〜24行目に適用）
          if rokuyo == "友":
            ws.cell(row=9, column=col_idx).value = "友"
            tomobiki_days.append(day)
            for r in range(7, 25):  # 7〜24行目を友引の色（薄緑）で着色
              ws.cell(row=r, column=col_idx).fill = tomobiki_fill
          elif next_rokuyo == "友":
            ws.cell(row=9, column=col_idx).value = "前"
            for r in range(7, 25):  # 7〜24行目を前日の色（1段薄い緑）で着色
              ws.cell(row=r, column=col_idx).fill = zenjitsu_fill
          else:
            ws.cell(row=9, column=col_idx).value = None

          # 日曜から土曜を1週間とするための区切り罫線（5〜24行目まで太線を統一適用）
          if dt.weekday() == 6:  # 6 = 日曜日
            for r in range(5, 25):  # 24行目まで太さを統一
              cell = ws.cell(row=r, column=col_idx)
              cell.border = Border(
                  left=thick_side,
                  right=cell.border.right,
                  top=cell.border.top,
                  bottom=cell.border.bottom,
                  diagonal=cell.border.diagonal,
                  diagonal_direction=cell.border.diagonal_direction,
                  outline=cell.border.outline,
                  vertical=cell.border.vertical,
                  horizontal=cell.border.horizontal,
              )
              if day > 1:
                prev_cell = ws.cell(row=r, column=col_idx - 1)
                prev_cell.border = Border(
                    left=prev_cell.border.left,
                    right=thick_side,
                    top=prev_cell.border.top,
                    bottom=prev_cell.border.bottom,
                    diagonal=prev_cell.border.diagonal,
                    diagonal_direction=prev_cell.border.diagonal_direction,
                    outline=prev_cell.border.outline,
                    vertical=prev_cell.border.vertical,
                    horizontal=prev_cell.border.horizontal,
                )

        else:
          ws.cell(row=7, column=col_idx).value = None
          ws.cell(row=8, column=col_idx).value = None
          ws.cell(row=9, column=col_idx).value = None

      # 集計数式の更新 (休 + 健 を公休としてカウント)
      for day in range(1, num_days + 1):
        col_idx = 3 + day
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.cell(
            row=5, column=col_idx
        ).value = (
            f'=COUNTIFS({col_letter}10:{col_letter}18,"休")+COUNTIFS({col_letter}10:{col_letter}18,"健")'
        )

      # 計算用データ構造の作成
      shift_matrix = {
          row["氏名"]: {d: "" for d in days}
          for _, row in edited_staff.iterrows()
      }
      daily_shain_off = {d: 0 for d in days}
      daily_total_off = {d: 0 for d in days}

      staff_info = edited_staff.set_index("氏名").to_dict(orient="index")

      late_shift_counts = {
          row["氏名"]: 0
          for _, row in edited_staff.iterrows()
          if row["遅番対応"]
      }

      # WEB画面（タブ2）で入力された事前予定（希望休・希望出勤・有給・健診）の読込
      for _, row in edited_pre_df.iterrows():
        name = row["氏名"]
        if name in shift_matrix:
          for d in days:
            val = str(row[f"{d}日"]).strip() if pd.notna(row[f"{d}日"]) else ""
            if val in ["休", "有", "健", "出"]:
              shift_matrix[name][d] = val
              if val in ["休", "健"]:  # 健診も公休扱い
                daily_total_off[d] += 1
                if staff_info[name]["区分"] == "社員":
                  daily_shain_off[d] += 1

      random.seed(45)

      # 1. 遅番の自動バランス割り当て (希望出勤「出」や「休」の日は回避)
      late_eligible = edited_staff[edited_staff["遅番対応"]]["氏名"].tolist()
      for d in days:
        if late_eligible:
          valid_late_candidates = []
          for m in late_eligible:
            if shift_matrix[m][d] != "":
              continue
            if d + 1 <= num_days:
              if shift_matrix[m][d + 1] != "":
                continue
              is_shain = staff_info[m]["区分"] == "社員"
              if is_shain and daily_shain_off[d + 1] >= 2:
                continue
              grp = staff_info[m]["担当顧客グループ"]
              grp_members = edited_staff[
                  edited_staff["担当顧客グループ"] == grp
              ]["氏名"].tolist()
              if grp != "共通" and any(
                  shift_matrix[gm][d + 1] in ["休", "健"]
                  for gm in grp_members
                  if gm != m
              ):
                continue
            valid_late_candidates.append(m)

          if valid_late_candidates:
            selected = min(
                valid_late_candidates, key=lambda x: late_shift_counts[x]
            )
            shift_matrix[selected][d] = "遅"
            late_shift_counts[selected] += 1

            if d + 1 <= num_days:
              shift_matrix[selected][d + 1] = "休"
              daily_total_off[d + 1] += 1
              if staff_info[selected]["区分"] == "社員":
                daily_shain_off[d + 1] += 1

      # 2. 残りの公休「休」の補完計算 (最大5連勤制限・グループ重複防止)
      sorted_staff = edited_staff.sort_values(by="区分", ascending=True)

      for _, staff in sorted_staff.iterrows():
        name = staff["氏名"]
        target_off = staff["公休数"]
        is_shain = staff["区分"] == "社員"
        grp = staff["担当顧客グループ"]

        current_offs = [
            d for d in days if shift_matrix[name][d] in ["休", "健"]
        ]
        needed_off = target_off - len(current_offs)

        for _ in range(needed_off):
          candidates = []
          for d in days:
            if shift_matrix[name][d] != "":
              continue

            if is_shain and daily_shain_off[d] >= 2:
              continue

            grp_members = edited_staff[
                edited_staff["担当顧客グループ"] == grp
            ]["氏名"].tolist()
            if grp != "共通" and any(
                shift_matrix[m][d] in ["休", "健"]
                for m in grp_members
                if m != name
            ):
              continue

            candidates.append(d)

          if not candidates:
            candidates = [
                d
                for d in days
                if shift_matrix[name][d] == ""
                and (not is_shain or daily_shain_off[d] < 2)
            ]
          if not candidates:
            candidates = [d for d in days if shift_matrix[name][d] == ""]

          def score_day(d):
            score = 0
            left_work = 0
            curr = d - 1
            while curr >= 1 and shift_matrix[name][curr] not in [
                "休",
                "健",
                "有",
            ]:
              left_work += 1
              curr -= 1

            right_work = 0
            curr = d + 1
            while curr <= num_days and shift_matrix[name][curr] not in [
                "休",
                "健",
                "有",
            ]:
              right_work += 1
              curr += 1

            total_consec = left_work + 1 + right_work

            if total_consec >= 6:
              score += 500 * (total_consec - 5)
            elif total_consec == 5:
              score += 50
            elif total_consec == 4:
              score += 15

            score += (10 - daily_total_off[d]) * 2
            score += random.uniform(0, 1)
            return score

          candidates.sort(key=score_day, reverse=True)
          best_d = candidates[0]

          shift_matrix[name][best_d] = "休"
          daily_total_off[best_d] += 1
          if is_shain:
            daily_shain_off[best_d] += 1

      # 3. Excelワークシートへ書き込み (出勤「出」は原本フォーマットに従い空欄に変換)
      for r in range(10, 18):
        emp_name = ws.cell(row=r, column=3).value
        if emp_name in shift_matrix:
          for day in range(1, num_days + 1):
            val = shift_matrix[emp_name][day]
            ws.cell(row=r, column=3 + day).value = (
                val if val in ["休", "有", "健", "遅"] else None
            )

      # 画面プレビュー用のDataFrame (希望出勤「出」は画面上でも空欄として表示)
      preview_matrix = {}
      for name, schedule in shift_matrix.items():
        preview_matrix[name] = {
            d: (val if val in ["休", "有", "健", "遅"] else "")
            for d, val in schedule.items()
        }
      res_df = pd.DataFrame(preview_matrix).T
      res_df.columns = [f"{d}日" for d in days]

      st.success("🎉 シフトのたたき台が完成しました！")
      st.dataframe(res_df, use_container_width=True)

      excel_buffer = io.BytesIO()
      wb.save(excel_buffer)
      excel_buffer.seek(0)

      st.download_button(
          label="📥 完成したシフト表（Excel）をダウンロード",
          data=excel_buffer,
          file_name=f"シフト（自動化）_{target_year}年{target_month}月.xlsx",
          mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          type="primary",
      )
